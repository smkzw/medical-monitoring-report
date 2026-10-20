#!/usr/bin/env python3
"""Write sponsor MM queries into a copy of the source EDC listing workbook ("首列版").

Each sheet gets a new first column 「申办方MM query」 (row label_row = label, row var_row = MMQUERY);
every query is written on its evidence row(s) (sheet,row from queries.json `ev`), several queries on one row are
joined with a blank line and prefixed by their id. Row numbers are unchanged because only a column is inserted.

python3 export_queries.py --norm work/norm --queries out/analysis/queries.json --out out [--wb edc]
"""
import argparse
import json
from collections import defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

FILL = PatternFill('solid', fgColor='FFF1D6')
HEAD = PatternFill('solid', fgColor='FF9900')


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--norm', required=True)
    ap.add_argument('--queries', required=True)
    ap.add_argument('--out', required=True)
    ap.add_argument('--wb', default=None, help='workbook alias to annotate (default: the one holding most evidence rows)')
    a = ap.parse_args()
    src = json.loads(Path(a.norm, '_source.json').read_text(encoding='utf-8'))
    Q = json.loads(Path(a.queries).read_text(encoding='utf-8'))
    sheets = src.get('sheets', {})
    by = defaultdict(lambda: defaultdict(list))
    wbcount = defaultdict(int)
    for q in Q:
        for e in q['ev']:
            sh, row = e[0], e[1]
            if sh in sheets and row:
                by[sh][row].append(q)
                wbcount[sheets[sh]['wb']] += 1
    alias = a.wb or (max(wbcount, key=wbcount.get) if wbcount else next(iter(src['workbooks'])))
    path = Path(src['workbooks'][alias])
    wb = openpyxl.load_workbook(path)
    lr, vr = src.get('label_row', 1), src.get('var_row', 2)
    n = 0
    for ws in wb.worksheets:
        ws.insert_cols(1)
        ws.cell(lr, 1, '申办方MM query').fill = HEAD
        ws.cell(lr, 1).font = Font(bold=True, color='FFFFFF')
        if vr:
            ws.cell(vr, 1, 'MMQUERY')
        ws.column_dimensions['A'].width = 60
        for row, qs in by.get(ws.title, {}).items():
            seen = []
            for q in qs:
                if q['id'] not in [s['id'] for s in seen]:
                    seen.append(q)
            c = ws.cell(row, 1, '\n\n'.join(f"[{q['id']}] {q['text']}" for q in seen))
            c.alignment = Alignment(wrap_text=True, vertical='top')
            c.fill = FILL
            n += len(seen)
        ws.freeze_panes = None
    out = Path(a.out) / f"{path.stem}_申办方MM_query_首列版.xlsx"
    wb.save(out)
    print(f'queries written: {n} cells-entries → {out}')


if __name__ == '__main__':
    main()
